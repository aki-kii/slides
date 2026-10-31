#!/usr/bin/env python3
"""11/5 CDK 登壇のスライドを .pptx で組む。Google スライドに取り込める形。

  python3 build_deck.py <出力先.pptx>

文字と箱は編集できる図形で置き、図だけ images/ の PNG を貼る。
仮の箇所には「仮」の札（図形名 kari）を付け、理由をノートに書く。
"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "images")
FONT = os.environ.get("DECK_FONT", "Noto Sans JP")
MONO = "Fira Code"

INK, SUB, LINE = "2c3e50", "6b7684", "c9d0d9"
GREEN, GREEN_BG = "0E8A6E", "e9f5f1"
RED, RED_BG = "DD344C", "fdecee"
FADE = "c9d0d9"
WHITE = "ffffff"
CODE_BG = "f5f7f9"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(hex_color):
    return RGBColor.from_string(hex_color.upper())


def text(slide, x, y, w, h, content, size=20, color=INK, bold=False, align="left", anchor="top", font=None, spacing=None):
    """content は文字列、または段落のリスト。段落は文字列か、(文字, 色, 太字) のリスト。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(box.text_frame, content, size, color, bold, align, anchor, font, spacing)
    return box


def fill_text(frame, content, size, color, bold, align, anchor, font, spacing=None):
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
        if spacing:
            p.line_spacing = spacing
        for run_text, run_color, run_bold in ([(para, color, bold)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = run_text
            r.font.size = Pt(size)
            r.font.bold = run_bold
            r.font.color.rgb = rgb(run_color)
            r.font.name = font or FONT
            r_pr = r._r.get_or_add_rPr()
            ea = etree.SubElement(r_pr, qn("a:ea"))
            ea.set("typeface", font or FONT)


def box(slide, x, y, w, h, line=LINE, fill=WHITE, content=None, size=18, color=INK, bold=True, radius=0.12, align="center", font=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = min(radius / min(w, h), 0.5)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(2)
    shape.shadow.inherit = False
    for ref in shape._element.xpath(".//a:effectRef"):
        ref.set("idx", "0")
    if content is not None:
        fill_text(shape.text_frame, content, size, color, bold, align, "middle", font)
        shape.text_frame.margin_left = shape.text_frame.margin_right = Inches(0.12)
    return shape


def arrow(slide, x1, y1, x2, y2, color=INK, width=2.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    ln = line.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return line


def new_slide(chapter=None, title=None, sub=None, kari=None, notes=None):
    s = prs.slides.add_slide(BLANK)
    if chapter:
        text(s, 0.6, 0.28, 8, 0.3, chapter, 13, GREEN, True)
    if title:
        text(s, 0.6, 0.55, 12.1, 0.6, title, 28, INK, True)
    if sub:
        text(s, 0.6, 1.14, 12.1, 0.6, sub, 16, SUB)
    if kari:
        width = 0.7 + 0.16 * len(kari)
        tag = box(s, 13.333 - 0.35 - width, 0.2, width, 0.32, RED, RED_BG, "仮：" + kari, 11, RED, True, 0.16)
        tag.name = "kari"
        tag.text_frame.word_wrap = False
    all_notes = []
    if kari:
        all_notes.append("仮: " + kari)
    if notes:
        all_notes.append(notes)
    if all_notes:
        s.notes_slide.notes_text_frame.text = "\n".join(all_notes)
    return s


def figure(slide, name, top=1.75, height=5.45, max_width=12.3):
    path = os.path.join(IMG, name + ".png")
    pic = slide.shapes.add_picture(path, 0, Inches(top), height=Inches(height))
    if pic.width > Inches(max_width):
        ratio = Inches(max_width) / pic.width
        pic.width, pic.height = Inches(max_width), int(pic.height * ratio)
    pic.left = int((prs.slide_width - pic.width) / 2)
    return pic


def code(slide, x, y, w, h, lines):
    shape = box(slide, x, y, w, h, LINE, CODE_BG, None)
    frame = shape.text_frame
    fill_text(frame, lines, 15, INK, False, "left", "middle", MONO)
    frame.margin_left = frame.margin_right = Inches(0.25)
    return shape


DEF_LINES = ("AWS Cloud Development Kit (AWS CDK) は、", "コードでクラウドインフラストラクチャを定義し、", "AWS CloudFormation を通じてプロビジョニングするための、", "オープンソースの", "ソフトウェア開発フレームワークです。")

# ---------- 導入 ----------
s = new_slide(kari="発表タイトルは未定", notes="発表タイトル、肩書きの書き方は未決。")
text(s, 1, 1.9, 11.3, 0.5, "JAWS-UG CDK支部 × 東京支部コラボ cdkdハンズオン", 20, SUB, align="center")
text(s, 1, 2.6, 11.3, 1.1, "AWS CDK とは？", 48, INK, True, "center")
text(s, 1, 4.3, 11.3, 0.9, ["2026.11.05（木）", "池田 晃尚（@akikii__）"], 20, SUB, align="center")

s = new_slide(title="自己紹介", kari="入れるかは未定", notes="#93 のスライドと同じ中身。入れるかは未決。")
s.shapes.add_picture(os.path.join(IMG, "me.jpeg"), Inches(1.2), Inches(1.9), height=Inches(3.4))
text(s, 5.2, 2.0, 7.4, 0.6, "アキキー | 池田 晃尚", 30, INK, True)
text(s, 5.2, 3.0, 7.4, 2.4, ["株式会社メイツ", "バックエンドエンジニア / SRE", "", "推しサービス: AWS CDK", "JAWS-UG CDK支部", "AWS Community Builders（Dev Tools, 2026〜）"], 20, INK)

s = new_slide("導入", "AWS CDK とは？", notes="公式の定義: https://docs.aws.amazon.com/ja_jp/cdk/v2/guide/home.html （根拠 10-30 01:30）")
text(s, 1.2, 2.1, 11.2, 3, [DEF_LINES[0], DEF_LINES[1], DEF_LINES[2], DEF_LINES[3] + DEF_LINES[4]], 28, INK, spacing=1.35)
text(s, 1.2, 6.2, 10.9, 0.4, "AWS CDK v2 デベロッパーガイドより", 14, SUB)

# ---------- IaC ----------
s = new_slide("IaC", "AWS CDK とは？", notes="根拠 10-30 01:45。")
text(s, 1.2, 2.1, 11.2, 3, [[(DEF_LINES[0], FADE, False)], [(DEF_LINES[1], GREEN, True)], [(DEF_LINES[2], FADE, False)], [(DEF_LINES[3], FADE, False), (DEF_LINES[4], GREEN, True)]], 28, spacing=1.35)

s = new_slide("IaC", notes="抜き出しは 10-31 09:27 の形。原文の言葉を変えず、つなぎだけ直している。")
text(s, 1.2, 2.2, 10.9, 1.8, "コードでクラウドインフラストラクチャを定義する、ソフトウェア開発フレームワークです", 32, INK, True, "center", spacing=1.3)
text(s, 1.2, 4.6, 10.9, 0.8, "IaC（Infrastructure as Code）の定義です！", 30, GREEN, True, "center")

s = new_slide("IaC", "定義する → 何を？", notes="根拠 10-30 01:53、01:54。宣言型の説明はここで一緒にする。手続き型は話さない（01:50）。")
text(s, 1.2, 2.5, 10.9, 1, "あるべき姿を定義する", 44, GREEN, True, "center")
text(s, 1.2, 4.2, 10.9, 0.6, "宣言型: どう作るかではなく、どうなっていてほしいかを書く", 22, INK, align="center")

s = new_slide("IaC", "プロビジョニングする → どうやって？", "あるべき姿との差分をデプロイする", "抜き出しの文に無い言葉", "10-31 09:27 の抜き出しには「プロビジョニング」が入っていない。この問いかけ（10-30 01:53）とのつながりは確認待ち。図は #93 の iac-diff を流用。")
figure(s, "iac-diff", 1.9, 5.0)

s = new_slide("IaC", "IaC の特徴", kari="置き場所は未定", notes="タイトルだけを見せて、中身は話さない（10-30 01:11）。置き場所は未決。題は #93 のスライドから。")
for i, t in enumerate(["再現性の高さ", "変更管理ができること", "構造を把握しやすい", "再利用のしやすさ"]):
    box(s, 1.2 + 5.6 * (i % 2), 2.2 + 1.8 * (i // 2), 5.3, 1.4, LINE, WHITE, t, 24)

# ---------- デプロイフロー ----------
s = new_slide("デプロイフロー", kari="抜き出しの文は確認待ち", notes="10-31 09:22 で「AWS CloudFormation を通じて」を戻してデプロイフローへつなぐと決めた。09:27 の抜き出しとどちらの文にするかは確認待ち。問いかけは 10-30 15:59。")
text(s, 1.2, 1.7, 11.2, 2.4, [[("コードでクラウドインフラストラクチャを定義し、", INK, False)], [("AWS CloudFormation を通じて", GREEN, True), ("プロビジョニングするための、", INK, False)], [("ソフトウェア開発フレームワークです", INK, False)]], 28, spacing=1.35)
text(s, 1.2, 4.8, 10.9, 0.8, "CDK は CloudFormation とどう関わってるの？", 30, INK, True, "center")

s = new_slide("デプロイフロー", "デプロイフロー", "cdk deploy の中で synth、publish、deploy が順に動く", notes="全体を先に見せる。ここから同じ図の一部を順に強調する。根拠 10-30 16:06、10-31 09:22。")
figure(s, "deploy-flow")

s = new_slide("デプロイフロー", "CDK Bootstrap", "デプロイに使うものを先に作っておく", "bootstrap を話すかは未定", "下書きの案A（bootstrap を話す）の中身。一覧のうち、ルックアップ用ロールと SSM パラメータは今の流れでは図に出てこない。")
figure(s, "deploy-flow-bootstrap")

s = new_slide("デプロイフロー", "synth（シンセサイズ）", "ソースコードをシンセサイズする", notes="「ビルド」とは書かない。")
figure(s, "deploy-flow-synth")

s = new_slide("デプロイフロー", "アセットグループ", "出力はテンプレートとそれ以外のファイル。いちばん大事なのはテンプレート。公式の呼び名は cloud assembly（cdk.out）", notes="「アセットグループ」は発表用の呼び名で、公式の用語ではない。マニフェストは目次の役割。ファイルの大きさでは manifest.json のほうがテンプレートより大きいことがある。synth の時点では zip も Docker イメージも作られない。")
figure(s, "deploy-flow-assets")

s = new_slide("デプロイフロー", "パブリッシュ", "アセットとして S3 バケットと ECR リポジトリへ上げる。上げるのはテンプレートとアセットで、マニフェストは上げない", notes="*.assets.json の一覧で確かめた。テンプレートも S3 へ上げるファイルとして載っている。Lambda のイメージの ECR push は軽く触れるだけ（下書きの口頭メモ）。")
figure(s, "deploy-flow-publish")

s = new_slide("デプロイフロー", "デプロイ", "新しいテンプレートとスタックの状態を比べて変更セットを作り、確認してスタックに反映すると AWS リソースが作られる。実際のリソースを見に行くわけではない", notes="公式は「スタックと、送った変更を比べる」。実際のリソースまで比べるのは drift-aware change set（cdk deploy --revert-drift）のとき。https://docs.aws.amazon.com/AWSCloudFormation/latest/UserGuide/using-cfn-updating-stacks-changesets.html 。リソースの話はこの章でする（10-31 09:39）。ロールバックと承認ゲートは扱わない。")
figure(s, "deploy-flow-deploy")

# ---------- プログラミング言語とフレームワーク ----------
s = new_slide("プログラミング言語とフレームワーク", notes="定義文のこの 2 語だけを残す（10-30 16:16）。公式は「コードで」だが、スライドでは「プログラミング言語で」に書き換える（16:15）。")
text(s, 1.2, 2.5, 10.9, 1, "プログラミング言語で", 44, GREEN, True, "center")
text(s, 1.2, 3.8, 10.9, 1, "フレームワーク", 44, GREEN, True, "center")
text(s, 1.2, 6.2, 10.9, 0.4, "公式の定義では「コードで」", 14, SUB, align="center")

s = new_slide("プログラミング言語とフレームワーク", "ループとパラメータ化", kari="コードの例は仮", notes="根拠 10-31 09:32。公式のベストプラクティスは、CloudFormation の Conditions や Parameters より、言語の機能で合成時に決めるよう勧めている。ただし CloudFormation にも今は Fn::ForEach がある。https://docs.aws.amazon.com/cdk/v2/guide/best-practices.html")
code(s, 0.8, 1.6, 7.0, 3.4, ["const stages = ['dev', 'stg', 'prod'];", "", "for (const stage of stages) {", "  new ApiStack(app, `Api-${stage}`, {", "    memorySize: memoryOf(stage),", "  });", "}"])
text(s, 8.3, 1.9, 4.4, 2.8, ["同じものを環境の数だけ作る", "", "違いは引数で渡す"], 24, INK, True)
text(s, 0.8, 5.5, 11.9, 0.9, "CloudFormation にも Fn::ForEach はあるが、使える値や操作はプログラミング言語のほうがずっと多い", 18, SUB)

s = new_slide("プログラミング言語とフレームワーク", "カスタムリソース", kari="入れるかは未定", notes="10-31 09:39 で本人は迷い中。もとは CloudFormation の機能で、テンプレートだけでも使える。https://docs.aws.amazon.com/cdk/api/v2/docs/aws-cdk-lib.custom_resources-readme.html")
text(s, 1.2, 1.9, 10.9, 0.8, "CloudFormation にない操作を、デプロイのときに自分のコードで動かす", 26, INK, True)
box(s, 1.2, 3.2, 5.3, 1.9, LINE, WHITE, [[("AwsCustomResource", INK, True)], [("API を 1 回呼ぶだけなら", SUB, False)]], 22)
box(s, 6.8, 3.2, 5.3, 1.9, LINE, WHITE, [[("Provider framework", INK, True)], [("自分で処理を書くなら", SUB, False)]], 22)
text(s, 1.2, 5.6, 10.9, 0.5, "もとは CloudFormation の機能。CDK だと簡単に書ける", 18, SUB)

s = new_slide("プログラミング言語とフレームワーク", "AWS CDK のつくり", kari="足すかは確認待ち", notes="Construct Library、CDK CLI、多言語対応を足すかは確認待ち（10-31 09:41）。https://docs.aws.amazon.com/ja_jp/cdk/v2/guide/home.html")
box(s, 0.9, 2.0, 3.65, 2.5, LINE, WHITE, [[("Construct Library", INK, True)], [("AWS が用意した部品の集まり", SUB, False)]], 20)
box(s, 4.85, 2.0, 3.65, 2.5, LINE, WHITE, [[("CDK CLI", INK, True)], [("cdk deploy などのコマンド", SUB, False)]], 20)
box(s, 8.8, 2.0, 3.65, 2.5, LINE, WHITE, [[("6 つの言語", INK, True)], [("TypeScript、JavaScript、Python、Java、C#、Go", SUB, False)]], 20)

s = new_slide("プログラミング言語とフレームワーク", "コンストラクト", "レベルが上がるほど抽象化も上がる", notes="この章の主役。抽象化はコンストラクトと 1 つの話（10-31 09:39）。https://docs.aws.amazon.com/cdk/v2/guide/constructs.html")
levels = [("L1", "CloudFormation のリソースそのまま", "CfnBucket"), ("L2", "安全な初期値つき", "s3.Bucket"), ("L3", "複数のリソースの構成一式", ["ApplicationLoad", "BalancedFargateService"])]
for i, (name, desc, example) in enumerate(levels):
    x = 0.9 + 4.0 * i
    shape = box(s, x, 4.1 - 0.75 * i, 3.55, 2.3 + 0.75 * i, GREEN, GREEN_BG if i == 2 else WHITE, None)
    top = 4.1 - 0.75 * i
    text(s, x + 0.2, top + 0.15, 3.15, 0.6, name, 30, GREEN, True, "center")
    text(s, x + 0.2, top + 0.85, 3.15, 0.8, desc, 17, INK, True, "center")
    text(s, x + 0.2, top + 1.6, 3.15, 0.6, example, 13, SUB, False, "center", font=MONO)

s = new_slide("プログラミング言語とフレームワーク", "アプリもスタックもコンストラクト", notes="スタックとアプリはコンストラクトの説明の中で触れる（10-31 09:39）。https://docs.aws.amazon.com/cdk/v2/guide/apps.html")
box(s, 2.0, 1.6, 9.3, 4.9, GREEN, GREEN_BG, None)
text(s, 2.3, 1.75, 4, 0.5, "アプリ", 22, GREEN, True)
box(s, 2.6, 2.4, 8.1, 3.7, LINE, WHITE, None)
text(s, 2.9, 2.55, 4, 0.5, "スタック", 22, INK, True)
for i, t in enumerate(["コンストラクト", "コンストラクト", "コンストラクト"]):
    box(s, 3.0 + 2.5 * i, 3.45, 2.3, 1.9, LINE, WHITE, t, 16)

s = new_slide("プログラミング言語とフレームワーク", "条件分岐", kari="コードの例は仮", notes="フレームワークとしての特徴として話す（10-31 09:32）。")
code(s, 2.0, 1.8, 9.3, 3.0, ["if (stage === 'prod') {", "  new MonitoringStack(app, 'Monitoring');", "}"])
text(s, 2.0, 5.3, 9.3, 0.6, "環境ごとの違いを if でそのまま書ける", 24, INK, True, "center")

s = new_slide("プログラミング言語とフレームワーク", "好きな CDK コンストラクト", kari="口頭で話す 1〜2 個は未定", notes="10-30 16:29 では全部 AWS 公式としていたが、npm の deploy-time-build のリポジトリは tmokmss/deploy-time-build で、作者も tmokmss。個人の OSS と分かるように書いた。16:32 で同じ理由で外した 2 つと同様に、残すかは本人が決める。DockerImageAsset の説明は本人のメモより短くした（各リポジトリへのコピーは DockerImageAsset だけではできないため）。")
rows = [("DockerImageAsset", "ビルドしたコンテナイメージを CDK Toolkit の ECR リポジトリに置く"), ("S3Deployment", "静的サイトのデプロイをいい感じにして、CloudFront のキャッシュの無効化までしてくれる"), ("NAT インスタンスのコンストラクト", "NAT インスタンスを構築する"), ("DeployTimeBuild（個人の OSS）", "tmokmss さんのコンストラクト")]
for i, (name, desc) in enumerate(rows):
    y = 1.6 + 1.2 * i
    box(s, 0.9, y, 11.5, 1.05, LINE, WHITE, [[(name, INK, True)], [(desc, SUB, False)]], 17, align="left")

# ---------- オープンソース ----------
s = new_slide("オープンソース", "AWS CDK は OSS", notes="軽く流す（10-31 09:25）。")
steps = ["AWS CDK は OSS", "外からのコントリビュートを受け付けている", "良いコンストラクトが作られていく"]
for i, t in enumerate(steps):
    x = 0.8 + 4.2 * i
    box(s, x, 2.6, 3.4, 1.9, GREEN if i == 2 else LINE, GREEN_BG if i == 2 else WHITE, t, 20)
    if i < 2:
        arrow(s, x + 3.4, 3.55, x + 4.2, 3.55)

oss_official = [("BLEA", "安全なベースライン環境を作る CDK テンプレート"), ("cdk-nag", "ルールで CDK アプリを確かめる"), ("jsii", "CDK を TypeScript 以外の言語から使えるようにする"), ("cdk-from-cfn", "CloudFormation テンプレートを CDK のコードに変換する"), ("cdk8s", "Kubernetes のアプリをコードで定義する"), ("AWS Solutions Constructs", "検証済みの構成のコンストラクト集"), ("cdk-monitoring-constructs", "監視を簡単に作るコンストラクト集")]
oss_japan = [("cdkd（後藤さん）", "CloudFormation を通さず AWS SDK を直接呼ぶ CDK CLI"), ("cdk-real-drift（後藤さん）", "テンプレートに書いていないプロパティまでドリフトを検出する"), ("awscdk-lint（山梨さん）", "CDK 用のリントのプラグイン"), ("cdk-preflight（くらいやーさん）", "デプロイ時の失敗を synth の時点で見つける"), ("cfn-exec-policy（くらいやーさん）", "CloudFormation 実行ロールを最小権限にする"), ("cfn-idle-cost（くらいやーさん）", "アクセスがなくてもかかる月額を出す")]
oss_world = [("projen", "プロジェクトの設定をコードで管理する"), ("cdk-dia", "構成図を自動で作る"), ("awesome-cdk", "CDK 関連の OSS や記事のまとめ")]


def oss_list(slide, x, y, w, items, size=15, step=0.66):
    for i, (name, desc) in enumerate(items):
        text(slide, x, y + step * i, w, 0.3, name, size, INK, True, font=None)
        text(slide, x, y + step * i + 0.29, w, 0.3, desc, size - 2, SUB)


s = new_slide("オープンソース", "関連する OSS（公式側）", kari="口頭で話す 3 つは未定", notes="スライドには全部載せ、口頭では 3 つくらいに絞る（10-30 01:05）。一覧は 10-30 16:24。説明は各リポジトリの説明から書いた。")
oss_list(s, 0.9, 1.45, 11.5, oss_official, 17, 0.77)

s = new_slide("オープンソース", "関連する OSS（日本・海外）", kari="口頭で話す 3 つは未定", notes="一覧は 10-30 16:24。")
text(s, 0.9, 1.3, 6.5, 0.4, "日本", 18, GREEN, True)
oss_list(s, 0.9, 1.8, 6.7, oss_japan, 15, 0.8)
text(s, 8.1, 1.3, 4.6, 0.4, "海外", 18, GREEN, True)
oss_list(s, 8.1, 1.8, 4.6, oss_world, 15, 0.8)

# ---------- 締め ----------
s = new_slide("締め", kari="締め方と渡し方は未定", notes="ここは決まっていない。ブリーフの「締め=後藤さん cdkd へ渡す」をもとにした仮の文。cdkd の説明はワークショップの README から。")
text(s, 1.2, 2.3, 10.9, 1, "ここからは後藤さんの cdkd へ", 40, INK, True, "center")
text(s, 1.2, 3.8, 10.9, 0.7, "CloudFormation を通さず、AWS SDK を直接呼ぶ CDK CLI", 22, SUB, align="center")

out = sys.argv[1]
prs.save(out)
print(len(prs.slides), "slides")
